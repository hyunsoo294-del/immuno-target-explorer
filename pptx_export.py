# -*- coding: utf-8 -*-
"""PPT from the same Plotly charts and data shown in the Streamlit browser."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

import charts

NAVY = RGBColor(0x0B, 0x1F, 0x33)
TEAL = RGBColor(0x1A, 0xA6, 0xA6)
INK = RGBColor(0x1F, 0x29, 0x37)
SLATE = RGBColor(0x5B, 0x6C, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF7, 0xF5, 0xF2)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE6, 0xE2, 0xDC)

FONT = "Calibri"
DPI = 180
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.36)

FIGURE_NOTES = [
    "Fig 1. RNA grouped bar: Normal tissue (HPA/GTEx, teal) vs Tumor (TCGA median, coral). Axis = log2(expression + 1).",
    "Fig 2. RNA log2 fold = log2((Tumor + 0.25) / (Normal + 0.25)). Positive = higher in tumor.",
    "Fig 3. Protein IHC 0-3: matched normal organ vs tumor (HPA). Navy = normal, gold = tumor.",
    "Fig 4. IHC Mann-Whitney effect. Positive = tumor staining ranks higher than matched normal.",
    "Fig 5. TCGA tumor RNA box plot (same samples as tab 3). Axis = log2(RSEM + 1). X = cancer type.",
    "Fig 6. TCGA tumor median vs adjacent-normal mean, log2(RSEM + 1).",
    "Fig 7. HPA IHC patient counts: High / Medium / Low / Not detected (same as tab 4a).",
    "Fig 8. CPTAC median tumor protein log2 ratio (same as tab 4b). >0 = above study reference.",
]

FIGURE_LEGENDS: dict[int, list[tuple[str, str]]] = {
    1: [("#1AA6A6", "Normal (HPA/GTEx)"), ("#FF5A5F", "Tumor (TCGA)")],
    2: [("#FF5A5F", "Higher in tumor"), ("#1AA6A6", "Higher in normal")],
    3: [("#0B1F33", "Normal IHC 0-3"), ("#E3B23C", "Cancer IHC 0-3")],
    4: [("#FF5A5F", "Tumor IHC higher"), ("#1AA6A6", "Normal IHC higher")],
    5: [("#FFD6D8", "Tumor RNA samples"), ("#000000", "Median")],
    6: [("#FF5A5F", "Tumor median"), ("#1AA6A6", "Adjacent normal")],
    7: [("#5B6C7D", "Not detected"), ("#1AA6A6", "Low"), ("#E3B23C", "Medium"), ("#FF5A5F", "High")],
    8: [("#1AA6A6", "Median log2 ratio")],
}


def _as_df(value) -> pd.DataFrame:
    if isinstance(value, pd.DataFrame):
        return value
    return pd.DataFrame()


def _set_run(run, *, size: int, bold: bool = False, color: RGBColor = INK) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT


def _text(slide, left, top, width, height, text: str, *, size: int = 12, bold: bool = False, color: RGBColor = INK, align=PP_ALIGN.LEFT, wrap: bool = False) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = str(text)
    if p.runs:
        _set_run(p.runs[0], size=size, bold=bold, color=color)


def _notes_block(slide, left, top, width, height, lines: list[str]) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(2)
        if p.runs:
            _set_run(p.runs[0], size=9, color=INK)


def _canvas(slide) -> None:
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()


def _card(slide, left, top, width, height) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = BORDER
    card.line.width = Pt(0.75)


def _hex(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _legend_under(slide, left, top, width, items: list[tuple[str, str]]) -> None:
    x = left
    chip = Inches(0.13)
    gap = Inches(0.08)
    for color, label in items:
        sq = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top + Inches(0.04), chip, chip)
        sq.fill.solid()
        sq.fill.fore_color.rgb = _hex(color)
        sq.line.color.rgb = RGBColor(0, 0, 0)
        sq.line.width = Pt(0.6)
        text_w = Inches(1.50) if len(items) <= 3 else Inches(1.18)
        _text(slide, x + chip + Inches(0.05), top, text_w, Inches(0.22), label, size=9, bold=True, color=RGBColor(0, 0, 0))
        x = x + chip + Inches(0.05) + text_w + gap


def _place(slide, png: bytes | None, left, top, w_in: float, h_in: float, empty: str) -> None:
    if png:
        slide.shapes.add_picture(io.BytesIO(png), left, top, width=Inches(w_in), height=Inches(h_in))
        return
    _text(slide, left + Inches(0.2), top + Inches(h_in / 2 - 0.15), Inches(w_in - 0.3), Inches(0.3), empty, size=11, color=SLATE)
    if png:
        slide.shapes.add_picture(io.BytesIO(png), left, top, width=Inches(w_in), height=Inches(h_in))
        return
    _text(slide, left + Inches(0.2), top + Inches(h_in / 2 - 0.15), Inches(w_in - 0.3), Inches(0.3), empty, size=11, color=SLATE)


def _short(value: Any, n: int = 20) -> str:
    text = " ".join(str(value or "").split())
    return text if len(text) <= n else text[: n - 1] + "."


def _ihc_view(bundle: dict) -> pd.DataFrame:
    analytics = _as_df(bundle.get("analytics_df"))
    if not analytics.empty:
        out = analytics.copy()
        if "protein_ihc_normal_0_3" not in out.columns:
            comparison = _as_df(bundle.get("comparison_df"))
            if not comparison.empty:
                merge_cols = ["cancer_type", "protein_ihc_normal_0_3", "protein_ihc_cancer_0_3", "ihc_mwu_pvalue", "ihc_mwu_effect", "ihc_mwu_significant"]
                avail = [c for c in merge_cols if c in comparison.columns]
                if len(avail) > 1:
                    out = out.merge(comparison[avail], on="cancer_type", how="left")
        return out
    comparison = _as_df(bundle.get("comparison_df"))
    if comparison.empty:
        return pd.DataFrame()
    rows = []
    for rec in comparison.to_dict("records"):
        if pd.notna(rec.get("protein_ihc_cancer_0_3")):
            rows.append(
                {
                    "cancer_type": rec.get("cancer_type"),
                    "ihc_weighted_0_3": rec.get("protein_ihc_cancer_0_3"),
                    "protein_ihc_normal_0_3": rec.get("protein_ihc_normal_0_3"),
                    "protein_ihc_cancer_0_3": rec.get("protein_ihc_cancer_0_3"),
                    "ihc_mwu_pvalue": rec.get("ihc_mwu_pvalue"),
                    "ihc_mwu_effect": rec.get("ihc_mwu_effect"),
                    "ihc_mwu_significant": rec.get("ihc_mwu_significant"),
                }
            )
    return pd.DataFrame(rows)


def _mpl():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Calibri", "Segoe UI", "Arial"],
            "axes.titlesize": 10,
            "axes.labelsize": 10,
            "axes.labelcolor": "#000000",
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "xtick.color": "#000000",
            "ytick.color": "#000000",
            "text.color": "#000000",
        }
    )
    return plt


def _layout_title(fig) -> str:
    try:
        return str(fig.layout.title.text or "")
    except Exception:
        return ""


def _layout_xlabel(fig) -> str:
    try:
        return str(fig.layout.xaxis.title.text or "")
    except Exception:
        return ""


def _layout_ylabel(fig) -> str:
    try:
        return str(fig.layout.yaxis.title.text or "")
    except Exception:
        return ""


def _marker_color(trace, fallback: str):
    try:
        color = trace.marker.color
    except Exception:
        return fallback
    if color is None:
        return fallback
    if isinstance(color, str) and color:
        return color
    if isinstance(color, (list, tuple)) and color:
        first = color[0]
        if isinstance(first, str):
            return list(color)
        return fallback
    return fallback


def _as_list(value) -> list:
    if value is None:
        return []
    return [v for v in list(value)]


def _style_axes(ax, *, box: bool = False) -> None:
    ax.tick_params(axis="both", colors="black", labelcolor="black", labelsize=10, width=1.0, length=3.5, pad=3)
    for axis in (ax.xaxis, ax.yaxis):
        axis.label.set_color("black")
        axis.label.set_fontsize(11)
        axis.label.set_fontweight("bold")
        axis.set_tick_params(labelcolor="black", color="black")
    for lab in list(ax.get_xticklabels()) + list(ax.get_yticklabels()):
        lab.set_color("black")
        lab.set_fontweight("bold")
        lab.set_fontsize(9)
    if box:
        for lab in ax.get_xticklabels():
            lab.set_rotation(55)
            lab.set_ha("right")
            lab.set_rotation_mode("anchor")
            lab.set_fontsize(9)
            lab.set_color("black")
            lab.set_fontweight("bold")


def _fig_to_png(fig, width_in: float, height_in: float) -> bytes | None:
    if fig is None:
        return None
    traces = list(getattr(fig, "data", []) or [])
    if not traces:
        return None
    plt = _mpl()
    fig_mpl, ax = plt.subplots(figsize=(width_in, height_in), dpi=DPI)
    fig_mpl.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    types = [getattr(t, "type", "") for t in traces]
    title = _layout_title(fig)
    xlabel = _layout_xlabel(fig)
    ylabel = _layout_ylabel(fig)
    is_box = "box" in types

    try:
        if is_box:
            _draw_box(ax, traces, fig)
            xlabel = xlabel or "Cancer type"
            ylabel = ylabel or "log2(RSEM + 1)"
        elif all(t == "bar" for t in types):
            _draw_bars(ax, traces, fig)
        else:
            plt.close(fig_mpl)
            return None
    except Exception:
        plt.close(fig_mpl)
        return None

    ax.set_title(title, loc="left", color="black", fontweight="bold", pad=6)
    if xlabel and "HPA tumor TMA" in xlabel:
        xlabel = "Number of patients"
    if xlabel:
        ax.set_xlabel(xlabel)
    if ylabel:
        ax.set_ylabel(ylabel)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#000000")
    ax.spines["bottom"].set_color("#000000")
    ax.grid(axis="x" if not is_box else "y", color="#E6E2DC", linewidth=0.6)
    if is_box:
        ax.grid(axis="x", visible=False)
    _style_axes(ax, box=is_box)
    if is_box:
        fig_mpl.subplots_adjust(left=0.12, right=0.98, top=0.88, bottom=0.42)
    else:
        is_stack = str(getattr(fig.layout, "barmode", None) or "") == "stack"
        fig_mpl.subplots_adjust(left=0.30 if is_stack else 0.26, right=0.97, top=0.88, bottom=0.18)
    buf = io.BytesIO()
    fig_mpl.savefig(buf, format="png", dpi=DPI, facecolor="#FFFFFF", pad_inches=0)
    plt.close(fig_mpl)
    return buf.getvalue()


def _draw_bars(ax, traces, fig) -> None:
    first = traces[0]
    labels = [_short(v, 18) for v in _as_list(first.y)]
    if not labels:
        raise ValueError("no labels")
    y = np.arange(len(labels))
    barmode = str(getattr(fig.layout, "barmode", None) or "group")
    n = len(traces)
    if barmode == "stack":
        left = np.zeros(len(labels))
        for trace in traces:
            vals = pd.to_numeric(pd.Series(_as_list(trace.x)), errors="coerce").fillna(0.0).to_numpy()
            if len(vals) != len(labels):
                vals = np.resize(vals, len(labels))
            color = _marker_color(trace, "#1AA6A6")
            if isinstance(color, list):
                color = color[0]
            ax.barh(y, vals, left=left, height=0.64, color=color, edgecolor="none")
            left = left + vals
    elif n > 1:
        bar_h = 0.72 / n
        for i, trace in enumerate(traces):
            vals = pd.to_numeric(pd.Series(_as_list(trace.x)), errors="coerce").fillna(0.0).to_numpy()
            if len(vals) != len(labels):
                vals = np.resize(vals, len(labels))
            offset = (i - (n - 1) / 2) * bar_h
            color = _marker_color(trace, "#1AA6A6" if i == 0 else "#FF5A5F")
            if isinstance(color, list):
                color = color[0]
            ax.barh(y + offset, vals, height=bar_h * 0.9, color=color, edgecolor="none")
    else:
        vals = pd.to_numeric(pd.Series(_as_list(first.x)), errors="coerce").fillna(0.0)
        color = _marker_color(first, "#FF5A5F")
        ax.barh(y, vals.tolist(), height=0.62, color=color, edgecolor="none")
    ax.set_yticks(y)
    ax.set_yticklabels(labels, color="#000000", fontweight="bold")
    ax.axvline(0, color="#000000", linewidth=0.7)
    ax.tick_params(axis="y", labelsize=8 if len(labels) > 12 else 9, pad=2, colors="#000000")


def _draw_box(ax, traces, fig) -> None:
    trace = next(t for t in traces if getattr(t, "type", "") == "box")
    xs = [str(v) for v in _as_list(trace.x)]
    ys = pd.to_numeric(pd.Series(_as_list(trace.y)), errors="coerce")
    work = pd.DataFrame({"x": xs, "y": ys}).dropna()
    if work.empty:
        raise ValueError("empty box")
    try:
        order = [str(v) for v in list(fig.layout.xaxis.categoryarray or [])]
    except Exception:
        order = []
    if not order:
        order = work.groupby("x")["y"].median().sort_values(ascending=False).index.tolist()
    data = [work.loc[work["x"] == name, "y"].to_numpy() for name in order]
    data = [d if len(d) else np.array([np.nan]) for d in data]
    tick_labels = [_short(v, 16) for v in order]
    positions = list(range(1, len(order) + 1))
    try:
        bp = ax.boxplot(data, positions=positions, tick_labels=tick_labels, patch_artist=True, flierprops={"markersize": 3, "markeredgecolor": "#FF5A5F"})
    except TypeError:
        bp = ax.boxplot(data, positions=positions, labels=tick_labels, patch_artist=True, flierprops={"markersize": 3, "markeredgecolor": "#FF5A5F"})
    for patch in bp["boxes"]:
        patch.set_facecolor("#FFD6D8")
        patch.set_edgecolor("#000000")
        patch.set_linewidth(0.8)
    for median in bp["medians"]:
        median.set_color("#000000")
        median.set_linewidth(1.2)
    ax.set_xticks(positions)
    ax.set_xticklabels(tick_labels, rotation=55, ha="right", rotation_mode="anchor", color="#000000", fontsize=8, fontweight="bold")
    ax.set_xlim(0.4, len(order) + 0.6)


def _header(slide, kicker: str, title: str, subtitle: str) -> None:
    _text(slide, MARGIN, Inches(0.12), Inches(12.5), Inches(0.20), kicker, size=9, bold=True, color=TEAL)
    _text(slide, MARGIN, Inches(0.30), Inches(12.5), Inches(0.28), title, size=16, bold=True, color=NAVY)
    _text(slide, MARGIN, Inches(0.56), Inches(12.5), Inches(0.20), subtitle, size=10, color=SLATE)


def _two_chart_slide(prs, *, kicker: str, title: str, subtitle: str, left_png, right_png, left_empty: str, right_empty: str, fig_left: str, fig_right: str, legend_left: list[tuple[str, str]], legend_right: list[tuple[str, str]]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _canvas(slide)
    _header(slide, kicker, title, subtitle)
    w, h = 6.20, 5.55
    _text(slide, MARGIN, Inches(0.76), Inches(6.2), Inches(0.18), fig_left, size=10, bold=True, color=RGBColor(0, 0, 0))
    _text(slide, Inches(6.78), Inches(0.76), Inches(6.2), Inches(0.18), fig_right, size=10, bold=True, color=RGBColor(0, 0, 0))
    _card(slide, MARGIN, Inches(0.96), Inches(6.32), Inches(6.08))
    _card(slide, Inches(6.78), Inches(0.96), Inches(6.32), Inches(6.08))
    _place(slide, left_png, Inches(0.42), Inches(1.02), w, h, left_empty)
    _place(slide, right_png, Inches(6.84), Inches(1.02), w, h, right_empty)
    _legend_under(slide, Inches(0.48), Inches(6.62), Inches(6.10), legend_left)
    _legend_under(slide, Inches(6.90), Inches(6.62), Inches(6.10), legend_right)
    return slide


def _footer(slide, symbol: str, page: int, total: int) -> None:
    _text(slide, MARGIN, Inches(7.24), Inches(10.8), Inches(0.18), f"{symbol}  |  Same data as the browser tabs  |  Research use only", size=8, color=SLATE)
    _text(slide, Inches(12.20), Inches(7.24), Inches(0.80), Inches(0.18), f"{page}/{total}", size=8, color=SLATE, align=PP_ALIGN.RIGHT)


def build_analysis_pptx(bundle: dict[str, Any]) -> bytes:
    identity = bundle.get("identity") or {}
    symbol = str(identity.get("symbol") or "gene")
    ensembl = str(identity.get("ensembl_id") or "")
    retrieved = bundle.get("retrieved_at") or datetime.now(timezone.utc).strftime("%Y-%m-%d")

    comparison = _as_df(bundle.get("comparison_df"))
    protein = comparison.dropna(subset=["protein_ihc_cancer_0_3"], how="any").copy() if not comparison.empty else pd.DataFrame()
    if protein.empty and not comparison.empty:
        protein = comparison.dropna(subset=["protein_ihc_normal_0_3"], how="any").copy()
    tcga = _as_df(bundle.get("tcga_df"))
    tcga_summary = _as_df(bundle.get("tcga_summary_df"))
    ihc = _ihc_view(bundle)
    cptac_summary = _as_df(bundle.get("cptac_summary_df"))

    w, h = 6.20, 5.55
    rna_l = _fig_to_png(charts.grouped_rna_bar(comparison), w, h)
    rna_r = _fig_to_png(charts.fold_bar(comparison, "log2_fold_normalized", "RNA log2 fold induction", "log2((Tumor+0.25)/(Normal+0.25))"), w, h)
    prot_l = _fig_to_png(charts.protein_grouped_bar(protein), w, h)
    prot_r = _fig_to_png(charts.protein_fold_horizontal(protein), w, h)
    tcga_l = _fig_to_png(charts.tcga_box(tcga), w, h)
    tcga_r = _fig_to_png(charts.tcga_tumor_normal_bar(tcga_summary), w, h)

    small_w, small_h = 6.20, 3.05
    ihc_png = _fig_to_png(charts.ihc_stacked_bar(ihc), small_w, small_h)
    cptac_png = _fig_to_png(charts.cptac_median_bar(cptac_summary), small_w, small_h)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    pages = []

    pages.append(
        _two_chart_slide(
            prs,
            kicker="BROWSER TAB 1  |  RNA EXPRESSION",
            title=f"{symbol}   {ensembl}",
            subtitle=f"Same RNA table and charts as the browser  |  {retrieved}",
            left_png=rna_l,
            right_png=rna_r,
            left_empty="RNA grouped chart not in this session.",
            right_empty="RNA fold chart not in this session.",
            fig_left="Fig 1",
            fig_right="Fig 2",
            legend_left=FIGURE_LEGENDS[1],
            legend_right=FIGURE_LEGENDS[2],
        )
    )
    pages.append(
        _two_chart_slide(
            prs,
            kicker="BROWSER TAB 2  |  PROTEIN EXPRESSION",
            title=f"{symbol}   protein IHC 0-3",
            subtitle="Same HPA IHC comparison as the browser",
            left_png=prot_l,
            right_png=prot_r,
            left_empty="Protein IHC chart not in this session.",
            right_empty="IHC effect chart not in this session.",
            fig_left="Fig 3",
            fig_right="Fig 4",
            legend_left=FIGURE_LEGENDS[3],
            legend_right=FIGURE_LEGENDS[4],
        )
    )
    pages.append(
        _two_chart_slide(
            prs,
            kicker="BROWSER TAB 3  |  TCGA DISTRIBUTION",
            title=f"{symbol}   TCGA RNA",
            subtitle="Same TCGA samples and summary as the browser",
            left_png=tcga_l,
            right_png=tcga_r,
            left_empty="TCGA box plot not in this session.",
            right_empty="TCGA summary chart not in this session.",
            fig_left="Fig 5",
            fig_right="Fig 6",
            legend_left=FIGURE_LEGENDS[5],
            legend_right=FIGURE_LEGENDS[6],
        )
    )

    last = prs.slides.add_slide(prs.slide_layouts[6])
    _canvas(last)
    _header(last, "BROWSER TAB 4  |  IHC PATHOLOGY + CPTAC", f"{symbol}   HPA IHC and CPTAC protein", "Same tab 4a / 4b data as the browser. Figure notes below.")
    _text(last, MARGIN, Inches(0.76), Inches(6.2), Inches(0.18), "Fig 7", size=10, bold=True, color=RGBColor(0, 0, 0))
    _text(last, Inches(6.78), Inches(0.76), Inches(6.2), Inches(0.18), "Fig 8", size=10, bold=True, color=RGBColor(0, 0, 0))
    _card(last, MARGIN, Inches(0.96), Inches(6.32), Inches(3.42))
    _card(last, Inches(6.78), Inches(0.96), Inches(6.32), Inches(3.42))
    _place(last, ihc_png, Inches(0.42), Inches(1.00), small_w, small_h, "HPA IHC chart not in this session.")
    _place(last, cptac_png, Inches(6.84), Inches(1.00), small_w, small_h, "CPTAC chart not in this session.")
    _legend_under(last, Inches(0.48), Inches(4.10), Inches(6.10), FIGURE_LEGENDS[7])
    _legend_under(last, Inches(6.90), Inches(4.10), Inches(6.10), FIGURE_LEGENDS[8])
    _card(last, MARGIN, Inches(4.48), Inches(12.62), Inches(2.62))
    _text(last, Inches(0.50), Inches(4.56), Inches(12.2), Inches(0.22), "Figure notes", size=11, bold=True, color=TEAL)
    _notes_block(last, Inches(0.50), Inches(4.80), Inches(12.2), Inches(2.18), FIGURE_NOTES)
    pages.append(last)

    total = len(pages)
    for i, slide in enumerate(pages, start=1):
        _footer(slide, symbol, i, total)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
